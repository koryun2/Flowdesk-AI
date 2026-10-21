import hashlib
import ipaddress
import logging
import socket
from html.parser import HTMLParser
from io import BytesIO
from pathlib import Path
from urllib import error, request
from urllib.parse import urlparse
from uuid import UUID, uuid4

from django.db import transaction
from django.utils.text import slugify
from rest_framework import serializers

from core.files import MAX_UPLOAD_BYTES, image_mime, read_limited, safe_name, store_bytes

from .client import KnowledgeAIError, request_answer, request_embeddings
from .local_ai import (
    EMBEDDING_DIMENSIONS,
    INSUFFICIENT_ANSWER,
    LOCAL_ANSWER_MODEL,
    LOCAL_EMBEDDING_MODEL,
    SECTION_BOOST,
    answer_from_sources,
    chunk_document,
    detect_sections,
    embed_text,
    hybrid_score,
    is_complex_question,
    leaks_draft,
    section_answer,
    specific_tokens,
    tokenize,
)
from .models import KnowledgeChunk, KnowledgeDocument
from .vectors import nearest_chunk_ids, write_vectors


logger = logging.getLogger("flowdesk.api")

MAX_DOCUMENT_CHARS = 60_000
MAX_CHUNKS = 80
MIN_RELEVANCE = 0.15
MAX_CONTEXT_CHUNKS = 3
TEXT_EXTENSIONS = {".txt", ".md", ".markdown", ".csv", ".json"}
HTML_EXTENSIONS = {".html", ".htm"}


def index_document(document: KnowledgeDocument, *, remote: bool = True) -> KnowledgeDocument:
    document.status = KnowledgeDocument.Status.PROCESSING
    document.error_message = ""
    document.save(update_fields=["status", "error_message", "updated_at"])
    try:
        pieces = chunk_document(document.content)[:MAX_CHUNKS]
        if not pieces:
            raise ValueError("The document has no readable text.")
        vectors, model_name = embed_document([content for _section, content in pieces], remote=remote)
        chunks = [
            KnowledgeChunk(
                id=uuid4(),
                organization=document.organization,
                document=document,
                position=position,
                section=section,
                content=content,
                embedding=vector,
                embedding_model=model_name,
            )
            for position, ((section, content), vector) in enumerate(zip(pieces, vectors))
        ]
        with transaction.atomic():
            document.chunks.all().delete()
            KnowledgeChunk.objects.bulk_create(chunks)
            document.status = KnowledgeDocument.Status.READY
            document.metadata = {
                **(document.metadata or {}),
                "characters": len(document.content),
                "chunk_count": len(chunks),
                "embedding_model": model_name,
            }
            document.error_message = ""
            document.save(update_fields=["status", "metadata", "error_message", "updated_at"])
        write_vectors(chunks)
    except (KnowledgeAIError, ValueError, TypeError) as exc:
        _mark_failed(document, str(exc))
    except Exception:
        logger.exception("Knowledge indexing failed for %s", document.pk)
        _mark_failed(document, "Indexing failed.")
    return document


def answer_question(organization, question: str, document_ids=None) -> dict:
    vector, model_name = embed_query(question)
    matches = ranked_chunks(organization, question, vector, model_name, document_ids)
    if not matches and model_name != LOCAL_EMBEDDING_MODEL:
        vector = embed_text(question)
        model_name = LOCAL_EMBEDDING_MODEL
        matches = ranked_chunks(organization, question, vector, model_name, document_ids)
    hints = detect_sections(question)
    if hints and len(hints) <= 3 and not is_complex_question(question):
        direct = direct_sections_answer(
            organization, question, hints, vector, model_name, document_ids
        )
        if direct:
            return direct
    contexts = context_sources(matches)
    sources = cited_sources(matches)
    if not sources:
        return {
            "answer": INSUFFICIENT_ANSWER,
            "sources": [],
            "model_name": LOCAL_ANSWER_MODEL,
        }
    try:
        answer, answer_model = request_answer(question, contexts)
    except KnowledgeAIError:
        answer, answer_model = answer_from_sources(question, contexts), LOCAL_ANSWER_MODEL
    if leaks_draft(answer):
        answer = answer_from_sources(question, contexts)
    return {"answer": answer, "sources": sources, "model_name": answer_model}


def prepare_document(organization, title: str, source_type: str, content: str, **fields) -> KnowledgeDocument:
    normalized = content.strip()
    if len(normalized) < 40:
        raise serializers.ValidationError(
            {"content": "Add at least 40 characters of readable text."}
        )
    if len(normalized) > MAX_DOCUMENT_CHARS:
        raise serializers.ValidationError(
            {"content": "This document is too long to index. Split it into smaller sources."}
        )
    checksum = hashlib.sha256(normalized.encode()).hexdigest()
    if KnowledgeDocument.objects.filter(organization=organization, checksum=checksum).exists():
        raise serializers.ValidationError(
            {"content": "This content is already in the knowledge base."}
        )
    document = KnowledgeDocument(
        organization=organization,
        title=title.strip(),
        slug=unique_slug(organization, title),
        source_type=source_type,
        content=normalized,
        checksum=checksum,
        status=KnowledgeDocument.Status.DRAFT,
        **fields,
    )
    document.full_clean()
    document.save()
    return document


def prepare_image_document(
    organization,
    title: str,
    data: bytes,
    *,
    file_name: str,
    mime_type: str,
    created_by=None,
) -> KnowledgeDocument:
    checksum = hashlib.sha256(data).hexdigest()
    if KnowledgeDocument.objects.filter(organization=organization, checksum=checksum).exists():
        raise serializers.ValidationError({"file": "This file is already in the knowledge base."})
    document = KnowledgeDocument(
        organization=organization,
        created_by=created_by,
        title=title.strip(),
        slug=unique_slug(organization, title),
        source_type=KnowledgeDocument.SourceType.FILE,
        content="",
        file_name=file_name,
        mime_type=mime_type,
        checksum=checksum,
        status=KnowledgeDocument.Status.READY,
        metadata={"characters": 0, "chunk_count": 0},
    )
    document.full_clean()
    document.save()
    store_bytes(document, "stored_file", data, file_name)
    return document


def embed_document(pieces: list[str], *, remote: bool) -> tuple[list[list[float]], str]:
    if remote:
        try:
            vectors, model_name = request_embeddings(pieces)
            _validate_vectors(vectors, len(pieces))
            return vectors, model_name
        except KnowledgeAIError as exc:
            logger.warning("Falling back to local embeddings: %s", exc)
    return [embed_text(piece) for piece in pieces], LOCAL_EMBEDDING_MODEL


def embed_query(question: str) -> tuple[list[float], str]:
    try:
        vectors, model_name = request_embeddings([question])
        _validate_vectors(vectors, 1)
        return vectors[0], model_name
    except KnowledgeAIError:
        return embed_text(question), LOCAL_EMBEDDING_MODEL


def context_sources(matches) -> list[dict]:
    return [
        {
            "document_id": str(chunk.document_id),
            "title": chunk.document.title,
            "excerpt": chunk.content[:500],
            "relevance": score,
        }
        for score, chunk in matches
    ]


def cited_sources(matches) -> list[dict]:
    sources = []
    for source in context_sources(matches):
        current = next((item for item in sources if item["document_id"] == source["document_id"]), None)
        if current is None:
            sources.append(dict(source))
            continue
        if source["excerpt"] not in current["excerpt"]:
            current["excerpt"] = f"{current['excerpt']}\n\n{source['excerpt']}"[:500]
    return sources


def selected_documents(organization, raw) -> list[UUID] | None:
    if raw in (None, ""):
        return None
    if not isinstance(raw, list):
        raise serializers.ValidationError({"document_ids": "Choose sources from this workspace."})
    if len(raw) > 20:
        raise serializers.ValidationError({"document_ids": "Choose 20 sources or fewer."})
    if not raw:
        return None
    ids = []
    for item in raw:
        try:
            ids.append(UUID(str(item)))
        except (TypeError, ValueError) as exc:
            raise serializers.ValidationError(
                {"document_ids": "Choose sources from this workspace."}
            ) from exc
    found = set(
        KnowledgeDocument.objects.filter(organization=organization, id__in=ids).values_list(
            "id", flat=True
        )
    )
    if found != set(ids):
        raise serializers.ValidationError({"document_ids": "Choose sources from this workspace."})
    return ids


def direct_sections_answer(organization, question, sections, vector, model_name, document_ids=None):
    answers = []
    sources = []
    for section in sections:
        direct = direct_section_answer(
            organization, question, section, vector, model_name, document_ids
        )
        if not direct:
            return None
        answers.append(direct["answer"])
        for source in direct["sources"]:
            current = next(
                (item for item in sources if item["document_id"] == source["document_id"]),
                None,
            )
            if current is None:
                sources.append(dict(source))
                continue
            if source["excerpt"] not in current["excerpt"]:
                current["excerpt"] = f"{current['excerpt']}\n\n{source['excerpt']}"[:500]
            current["relevance"] = max(current["relevance"], source["relevance"])
    return {
        "answer": "\n\n".join(answers),
        "sources": sources,
        "model_name": LOCAL_ANSWER_MODEL,
    }


def direct_section_answer(
    organization, question: str, section: str, vector: list[float], model_name: str, document_ids=None
):
    chunks = list(
        _chunk_base(organization, model_name, document_ids).filter(section=section)
    )
    if not chunks:
        return None
    grouped: dict = {}
    for chunk in chunks:
        grouped.setdefault(chunk.document_id, []).append(chunk)
    specific = specific_tokens(question, section)
    ranked = []
    for parts in grouped.values():
        document = parts[0].document
        overlap = len(specific & set(tokenize(document.content[:8000]))) if specific else 1
        best = 0.0
        for chunk in parts:
            embedding = chunk.embedding if isinstance(chunk.embedding, list) else []
            if len(embedding) == len(vector):
                best = max(best, hybrid_score(question, chunk.content, vector, embedding))
        ranked.append((overlap, best, document, parts))
    ranked.sort(key=lambda item: (item[0], item[1]), reverse=True)
    overlap, score, document, parts = ranked[0]
    if specific and overlap == 0:
        return None
    parts.sort(key=lambda chunk: chunk.position)
    answer = section_answer([chunk.content for chunk in parts])
    if not answer:
        return None
    return {
        "answer": answer,
        "sources": [
            {
                "document_id": str(document.id),
                "title": document.title,
                "excerpt": parts[0].content[:500],
                "relevance": min(1.0, round(score + SECTION_BOOST, 4)),
            }
        ],
        "model_name": LOCAL_ANSWER_MODEL,
    }


def ranked_chunks(organization, question: str, vector: list[float], model_name: str, document_ids=None):
    hints = set(detect_sections(question))
    base = _chunk_base(organization, model_name, document_ids)
    pool = {}
    candidate_ids = nearest_chunk_ids(organization.id, vector, model_name)
    candidates = base.filter(id__in=candidate_ids) if candidate_ids else base
    for chunk in candidates:
        pool[chunk.id] = chunk
    if hints:
        for chunk in base.filter(section__in=hints):
            pool[chunk.id] = chunk
    scored = []
    for chunk in pool.values():
        embedding = chunk.embedding
        if not isinstance(embedding, list) or len(embedding) != len(vector):
            continue
        score = hybrid_score(question, chunk.content, vector, embedding)
        if chunk.section in hints:
            score = min(1.0, round(score + SECTION_BOOST, 4))
        if score >= MIN_RELEVANCE:
            scored.append((score, chunk))
    best = {}
    for score, chunk in scored:
        key = (chunk.document_id, chunk.section or "")
        current = best.get(key)
        if current is None or score > current[0]:
            best[key] = (score, chunk)
    ranked = sorted(best.values(), key=lambda item: item[0], reverse=True)
    return ranked[:MAX_CONTEXT_CHUNKS]


def _chunk_base(organization, model_name: str, document_ids):
    base = KnowledgeChunk.objects.filter(
        organization=organization,
        embedding_model=model_name,
        document__status=KnowledgeDocument.Status.READY,
    ).select_related("document")
    if document_ids:
        base = base.filter(document_id__in=document_ids)
    return base


def store_document_content(document: KnowledgeDocument, content: str, **fields) -> bool:
    normalized = content.strip()
    if len(normalized) < 40:
        raise serializers.ValidationError({"content": "Add at least 40 characters of readable text."})
    if len(normalized) > MAX_DOCUMENT_CHARS:
        raise serializers.ValidationError(
            {"content": "This document is too long to index. Split it into smaller sources."}
        )
    checksum = hashlib.sha256(normalized.encode()).hexdigest()
    if (
        KnowledgeDocument.objects.filter(organization=document.organization, checksum=checksum)
        .exclude(pk=document.pk)
        .exists()
    ):
        raise serializers.ValidationError(
            {"content": "This content is already in the knowledge base."}
        )
    changed = checksum != document.checksum
    document.content = normalized
    document.checksum = checksum
    update_fields = ["content", "checksum", "updated_at"]
    for name, value in fields.items():
        setattr(document, name, value)
        update_fields.append(name)
    document.save(update_fields=update_fields)
    return changed


def replace_image_document(document: KnowledgeDocument, data: bytes, **fields) -> KnowledgeDocument:
    checksum = hashlib.sha256(data).hexdigest()
    if (
        KnowledgeDocument.objects.filter(organization=document.organization, checksum=checksum)
        .exclude(pk=document.pk)
        .exists()
    ):
        raise serializers.ValidationError({"file": "This file is already in the knowledge base."})
    document.chunks.all().delete()
    document.content = ""
    document.checksum = checksum
    document.status = KnowledgeDocument.Status.READY
    document.error_message = ""
    document.metadata = {"characters": 0, "chunk_count": 0}
    update_fields = [
        "content",
        "checksum",
        "status",
        "error_message",
        "metadata",
        "updated_at",
    ]
    for name, value in fields.items():
        setattr(document, name, value)
        update_fields.append(name)
    document.save(update_fields=update_fields)
    store_bytes(document, "stored_file", data, fields.get("file_name") or document.file_name)
    return document


def clear_stored_file(document: KnowledgeDocument) -> None:
    if not document.stored_file:
        return
    document.stored_file.delete(save=False)
    document.stored_file = ""
    document.save(update_fields=["stored_file", "updated_at"])


def rename_document(document: KnowledgeDocument, title: str) -> KnowledgeDocument:
    cleaned = title.strip()
    if not 3 <= len(cleaned) <= 240:
        raise serializers.ValidationError({"title": "Enter a title between 3 and 240 characters."})
    if cleaned == document.title:
        return document
    document.title = cleaned
    document.slug = unique_slug(document.organization, cleaned, exclude=document.pk)
    document.save(update_fields=["title", "slug", "updated_at"])
    return document


def unique_slug(organization, title: str, *, exclude=None) -> str:
    base = slugify(title)[:100] or "document"
    slug = base
    suffix = 2
    while (
        KnowledgeDocument.objects.filter(organization=organization, slug=slug)
        .exclude(pk=exclude)
        .exists()
    ):
        slug = f"{base[:90]}-{suffix}"
        suffix += 1
    return slug


def read_upload(upload) -> tuple[bytes, str, str, str]:
    if upload is None:
        raise serializers.ValidationError({"file": "Choose a file."})
    data = read_limited(upload)
    name = safe_name(getattr(upload, "name", "") or "document.txt")
    mime = image_mime(name)
    if mime:
        return data, name, mime, ""
    suffix = Path(name).suffix.lower()
    if suffix in TEXT_EXTENSIONS:
        text = data.decode("utf-8", errors="replace")
    elif suffix in HTML_EXTENSIONS:
        text = _html_text(data.decode("utf-8", errors="replace"))
    elif suffix == ".pdf":
        text = _pdf_text(data)
    elif suffix == ".docx":
        text = _docx_text(data)
    else:
        raise serializers.ValidationError(
            {"file": "Upload a .txt, .md, .html, .csv, .json, .pdf, .docx, or an image."}
        )
    content_type = getattr(upload, "content_type", "") or "application/octet-stream"
    return data, name, content_type, text


def fetch_url_text(url: str) -> str:
    public_url = _public_http_url(url)
    http_request = request.Request(
        public_url,
        headers={"User-Agent": "FlowdeskKnowledge/1.0"},
        method="GET",
    )
    try:
        with request.urlopen(http_request, timeout=10) as response:
            raw = response.read(MAX_UPLOAD_BYTES + 1)
            content_type = response.headers.get("Content-Type", "")
    except error.HTTPError as exc:
        raise serializers.ValidationError(
            {"source_url": f"The URL returned status {exc.code}."}
        ) from exc
    except (error.URLError, TimeoutError) as exc:
        raise serializers.ValidationError({"source_url": "The URL could not be reached."}) from exc
    if len(raw) > MAX_UPLOAD_BYTES:
        raise serializers.ValidationError({"source_url": "That page is too large to import."})
    text = raw.decode("utf-8", errors="replace")
    if "html" in content_type or text.lstrip().lower().startswith("<!doctype") or "<html" in text[:200].lower():
        text = _html_text(text)
    return text


def _validate_vectors(vectors, expected: int) -> None:
    if len(vectors) != expected or any(
        not isinstance(vector, list) or len(vector) != EMBEDDING_DIMENSIONS for vector in vectors
    ):
        raise KnowledgeAIError("AI service returned an unexpected embedding.")


def _mark_failed(document: KnowledgeDocument, message: str) -> None:
    document.status = KnowledgeDocument.Status.FAILED
    document.error_message = message[:1000]
    document.save(update_fields=["status", "error_message", "updated_at"])


def _public_http_url(url: str) -> str:
    parsed = urlparse((url or "").strip())
    host = parsed.hostname
    if parsed.scheme not in {"http", "https"} or not host:
        raise serializers.ValidationError({"source_url": "Enter an http or https URL."})
    if host == "localhost" or host.endswith(".local"):
        raise serializers.ValidationError({"source_url": "That URL cannot be imported."})
    port = parsed.port or (443 if parsed.scheme == "https" else 80)
    try:
        infos = socket.getaddrinfo(host, port)
    except socket.gaierror as exc:
        raise serializers.ValidationError({"source_url": "That URL could not be resolved."}) from exc
    for info in infos:
        address = ipaddress.ip_address(info[4][0])
        if not address.is_global:
            raise serializers.ValidationError({"source_url": "That URL cannot be imported."})
    return parsed.geturl()


def _html_text(markup: str) -> str:
    parser = _TextExtractor()
    parser.feed(markup)
    return parser.text()


def _pdf_text(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise serializers.ValidationError({"file": "PDF import is unavailable."}) from exc
    try:
        reader = PdfReader(BytesIO(data))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except Exception as exc:
        raise serializers.ValidationError({"file": "The PDF could not be read."}) from exc


def _docx_text(data: bytes) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise serializers.ValidationError({"file": "Word import is unavailable."}) from exc
    try:
        document = Document(BytesIO(data))
        return "\n".join(paragraph.text for paragraph in document.paragraphs)
    except Exception as exc:
        raise serializers.ValidationError({"file": "The Word document could not be read."}) from exc


class _TextExtractor(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts: list[str] = []
        self.skip = 0

    def handle_starttag(self, tag, attrs):
        if tag in {"script", "style", "noscript"}:
            self.skip += 1

    def handle_endtag(self, tag):
        if tag in {"script", "style", "noscript"} and self.skip:
            self.skip -= 1
        if tag in {"p", "div", "br", "li", "h1", "h2", "h3", "tr"}:
            self.parts.append("\n")

    def handle_data(self, data):
        if not self.skip:
            self.parts.append(data)

    def text(self) -> str:
        return "".join(self.parts)
